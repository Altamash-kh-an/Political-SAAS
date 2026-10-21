import csv
from io import BytesIO, StringIO
from pathlib import PurePath
import re
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


IMPORT_COLUMNS = (
    "booth_number",
    "polling_station_name",
    "polling_station_address",
    "block_name",
    "village_name",
    "panchayat_name",
)
MAX_IMPORT_ROWS = 5000


class BoothImportFormatError(ValueError):
    pass


def _text_value(value: object, column_name: str, number_format: str = "General") -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, str):
        return value
    if column_name == "booth_number" and isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        if isinstance(value, int) and re.fullmatch(r"0+", number_format):
            return str(value).zfill(len(number_format))
    return str(value)


def _validate_headers(header: list[object]) -> dict[str, int]:
    normalized = [
        value.strip() if isinstance(value, str) else ""
        for value in header
    ]
    duplicates = {
        name for name in normalized if name and normalized.count(name) > 1
    }
    if duplicates:
        raise BoothImportFormatError(
            f"Duplicate column headers: {', '.join(sorted(duplicates))}"
        )
    missing = [name for name in IMPORT_COLUMNS if name not in normalized]
    if missing:
        raise BoothImportFormatError(
            f"Missing required columns: {', '.join(missing)}"
        )
    return {name: normalized.index(name) for name in IMPORT_COLUMNS}


def parse_booth_import(filename: str, contents: bytes) -> list[dict[str, object]]:
    extension = PurePath(filename).suffix.lower()
    if extension == ".csv":
        return _parse_csv(contents)
    if extension == ".xlsx":
        return _parse_xlsx(contents)
    raise BoothImportFormatError("Choose an .xlsx or .csv file.")


def _parse_csv(contents: bytes) -> list[dict[str, object]]:
    try:
        text = contents.decode("utf-8-sig")
        reader = csv.reader(StringIO(text, newline=""))
        header = next(reader)
        column_indexes = _validate_headers(header)
        rows = []
        for row_number, values in enumerate(reader, start=2):
            if not any(value.strip() for value in values):
                continue
            if row_number > MAX_IMPORT_ROWS + 1:
                raise BoothImportFormatError(
                    f"Files may contain no more than {MAX_IMPORT_ROWS} data rows."
                )
            parsed_row = {
                "row_number": row_number,
                **{
                    name: values[index].strip() if index < len(values) else ""
                    for name, index in column_indexes.items()
                },
            }
            has_unlabeled_values = any(
                index >= len(header) or not header[index].strip()
                for index, value in enumerate(values)
                if value.strip()
            )
            if has_unlabeled_values:
                parsed_row["parse_errors"] = [
                    "This row has values under an unlabeled column."
                ]
            rows.append(parsed_row)
        return rows
    except UnicodeDecodeError as exc:
        raise BoothImportFormatError("CSV files must use UTF-8 encoding.") from exc
    except StopIteration as exc:
        raise BoothImportFormatError("The file is empty.") from exc
    except csv.Error as exc:
        raise BoothImportFormatError("The CSV file could not be parsed.") from exc


def _parse_xlsx(contents: bytes) -> list[dict[str, object]]:
    try:
        with ZipFile(BytesIO(contents)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 50 * 1024 * 1024:
                raise BoothImportFormatError(
                    "The expanded Excel file must not exceed 50 MB."
                )
        workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
    except BoothImportFormatError:
        raise
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise BoothImportFormatError("The Excel file could not be opened.") from exc

    try:
        worksheet = workbook.active
        if worksheet.max_column > 100:
            raise BoothImportFormatError(
                "Excel worksheets may contain no more than 100 columns."
            )
        row_iterator = worksheet.iter_rows()
        try:
            header_cells = next(row_iterator)
        except StopIteration as exc:
            raise BoothImportFormatError("The worksheet is empty.") from exc
        header_values = [cell.value for cell in header_cells]
        column_indexes = _validate_headers(header_values)
        rows = []
        for row_number, cells in enumerate(row_iterator, start=2):
            if not any(cell.value is not None and str(cell.value).strip() for cell in cells):
                continue
            if row_number > MAX_IMPORT_ROWS + 1:
                raise BoothImportFormatError(
                    f"Files may contain no more than {MAX_IMPORT_ROWS} data rows."
                )
            values = {}
            for name, index in column_indexes.items():
                if index >= len(cells):
                    values[name] = ""
                else:
                    cell = cells[index]
                    values[name] = _text_value(
                        cell.value,
                        name,
                        cell.number_format or "General",
                    ).strip()
            parsed_row = {"row_number": row_number, **values}
            if any(
                index >= len(header_values)
                or not isinstance(header_values[index], str)
                or not header_values[index].strip()
                for index, cell in enumerate(cells)
                if cell.value is not None and str(cell.value).strip()
            ):
                parsed_row["parse_errors"] = [
                    "This row has values under an unlabeled column."
                ]
            rows.append(parsed_row)
        return rows
    finally:
        workbook.close()
